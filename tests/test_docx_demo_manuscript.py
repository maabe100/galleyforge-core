"""Regression tests. All manuscripts here are fictional (no real people, institutions or data)."""
from pathlib import Path

import pytest
from docx import Document

from galleyforge_core.docx_parser_v1 import build_manuscript_dict_from_docx
from galleyforge_core.figures import attach_figure_files
from galleyforge_core.html_renderer import render_html
from galleyforge_core.parser import manuscript_to_schema_v1, parse_manuscript_file
from galleyforge_core.schema_v1 import ensure_schema_defaults
from galleyforge_core.validator_v1 import validate_manuscript_v1

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "base_template.html"
FICTIONAL_DOCX = ROOT / "examples" / "fictional_test_manuscript.docx"


@pytest.fixture(scope="module")
def fictional():
    return ensure_schema_defaults(build_manuscript_dict_from_docx(FICTIONAL_DOCX))


# ---- the fictional test manuscript (Word "Title" style, Unicode superscripts, table, figure caption) ----

def test_title_and_front_matter(fictional):
    assert fictional["title"].startswith("Effects of a Simulated Osmotic Challenge")
    assert len(fictional["front_matter"]) == 1 and "FICTIONAL DEMO MANUSCRIPT" in fictional["front_matter"][0]


def test_authors_and_affiliations(fictional):
    assert [(a["name"], a["affiliation_indices"]) for a in fictional["authors"]] == [
        ("Alex Example", [1]), ("Jamie Sample", [2]), ("Robin Placeholder", [1]),
    ]
    assert [a["text"] for a in fictional["affiliations"]] == [
        "Department of Synthetic Biology, Example Research Institute, Fictional City",
        "Laboratory of Model Systems, Sample University, Fictional City",
    ]


def test_correspondence_and_keywords(fictional):
    assert fictional["corresponding"] == {"author_index": 0, "email": "alex.example@example.org"}
    assert fictional["keywords"] == [
        "cell volume", "osmotic stress", "simulation", "workflow validation", "synthetic data",
    ]


def test_sections(fictional):
    s = fictional["sections"]
    for key in ("abstract", "introduction", "methods", "results", "discussion", "conclusion"):
        assert s[key], key
    assert s["ethics"].startswith("Not applicable")
    assert s["funding"] == "None." and s["disclosure_statement"] == "None."
    assert s["data_availability"].startswith("All illustrative summary values")
    assert s["author_contributions"].startswith("All named authors")


def test_table_and_figure_caption(fictional):
    (table,) = fictional["tables"]
    assert table["label"] == "Table 1" and table["caption"] == "Synthetic recovery index at 30 minutes"
    assert table["rows"] == [
        ["Condition", "Synthetic observations", "Mean index (unitless)"],
        ["Reference", "12", "78.2"],
        ["Simulated intervention", "12", "84.7"],
    ]
    (figure,) = fictional["figures"]
    assert figure["label"] == "Figure 1" and figure["caption"].startswith("Placeholder caption") and figure["file"] == ""
    # captions must not leak into the running text
    assert "Table 1." not in fictional["sections"]["results"].split("(Table 1)")[-1]
    assert "Placeholder caption" not in " ".join(str(v) for v in fictional["sections"].values())


def test_no_validation_errors(fictional):
    assert [i for i in validate_manuscript_v1(fictional) if i.severity == "error"] == []


def test_html_output(fictional):
    html = render_html(fictional, TEMPLATE)
    assert "{{" not in html
    assert "<title>Effects of a Simulated Osmotic Challenge" in html
    assert "<table>" in html and "<caption><strong>Table 1.</strong>" in html
    assert "<figcaption><strong>Figure 1.</strong>" in html
    assert "<strong>Correspondence:</strong> Alex Example (alex.example@example.org)" in html
    assert "<strong>Keywords:</strong>" in html and "FICTIONAL DEMO MANUSCRIPT" in html


# ---- other layouts ----

def _save(tmp_path, build):
    doc = Document()
    build(doc)
    path = tmp_path / "t.docx"
    doc.save(path)
    return build_manuscript_dict_from_docx(path)


def test_legacy_layout_double_space_authors(tmp_path):
    def build(doc):
        for text in [
            "Fictional Example: Structured Conversion of a Demonstration Manuscript",
            "Alex Example 1,2  Sam Sample 2",
            "1Department of Examples, Fictional University, Exampletown, Country",
            "2Institute of Samples, Demo University, Sampleville, Country",
            "Correspondence to: Alex Example, Department of Examples. E-mail: alex.example@example.org",
            "Keywords: demonstration; structured manuscript; HTML galley",
            "Abstract", "Invented abstract text.",
            "References", "1 Example A. A fictional reference. Journal of Examples. 2026;1:1-2.",
        ]:
            doc.add_paragraph(text)

    data = _save(tmp_path, build)
    assert [a["name"] for a in data["authors"]] == ["Alex Example", "Sam Sample"]
    assert data["corresponding"]["email"] == "alex.example@example.org"
    assert data["keywords"] == ["demonstration", "structured manuscript", "HTML galley"]
    assert data["sections"]["references"] == ["1 Example A. A fictional reference. Journal of Examples. 2026;1:1-2."]


def test_real_word_superscript_formatting(tmp_path):
    def build(doc):
        doc.add_paragraph("Probe", style="Title")
        p = doc.add_paragraph()
        for name, label, sep in (("Alex Example", "1", ", "), ("Jamie Sample", "2,3", "")):
            p.add_run(name)
            p.add_run(label).font.superscript = True
            p.add_run(sep)
        for label, text in (("1", "Department of Examples, Fictional University"),
                            ("2", "Institute of Samples, Demo University"),
                            ("3", "Laboratory of Tests, Example Institute")):
            q = doc.add_paragraph()
            q.add_run(label).font.superscript = True
            q.add_run(" " + text)
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Invented.")

    data = _save(tmp_path, build)
    assert [(a["name"], a["affiliation_indices"]) for a in data["authors"]] == [
        ("Alex Example", [1]), ("Jamie Sample", [2, 3]),
    ]
    assert len(data["affiliations"]) == 3


def test_table_caption_after_table(tmp_path):
    def build(doc):
        doc.add_paragraph("T", style="Title")
        table = doc.add_table(rows=2, cols=2)
        for r, row in enumerate((("a", "b"), ("1", "2"))):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
        doc.add_paragraph("Table 7. Caption below the table")

    data = _save(tmp_path, build)
    assert data["tables"][0]["label"] == "Table 7" and data["tables"][0]["caption"] == "Caption below the table"


def test_attach_figure_files_by_number(tmp_path):
    (tmp_path / "Fig_1.jpg").write_bytes(b"x")
    (tmp_path / "extra.png").write_bytes(b"x")
    result = attach_figure_files([{"label": "Figure 1", "caption": "c", "file": ""}], tmp_path)
    assert result[0]["file"] == "Fig_1.jpg"
    assert result[1]["file"] == "extra.png"


# ---- plain-text demo ----

def test_plain_text_demo_converts_to_schema_v1():
    data = manuscript_to_schema_v1(parse_manuscript_file(str(ROOT / "examples" / "manuscript.txt")))
    assert [a["name"] for a in data["authors"]] == ["Jane Doe", "John Smith"]
    assert set(data["sections"]) >= {"abstract", "introduction", "methods", "results", "discussion", "references"}
    assert data["sections"]["references"] == ["Example reference one.", "Example reference two."]
    html = render_html(ensure_schema_defaults(data), TEMPLATE)
    assert "{{" not in html and "<h2>Abstract</h2>" in html


def test_output_matches_json_schema(fictional):
    jsonschema = pytest.importorskip("jsonschema")
    import json

    schema = json.loads((ROOT / "galleyforge_core" / "manuscript_v1.schema.json").read_text(encoding="utf-8"))
    jsonschema.validate(fictional, schema)
